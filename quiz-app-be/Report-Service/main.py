"""
Report Service - Dịch vụ xuất báo cáo / bảng điểm
===================================================
Chạy độc lập trên port 8006
Kết nối trực tiếp MongoDB để lấy dữ liệu submissions, tests, students
Cung cấp API cho Task-Service (API Gateway) gọi vào
"""

import io
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import pandas as pd
import uvicorn
from bson import ObjectId
from dotenv import load_dotenv
from fastapi import FastAPI, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from motor.motor_asyncio import AsyncIOMotorClient

# ==========================================
# CẤU HÌNH
# ==========================================
load_dotenv(dotenv_path=Path(".") / ".env")

FAPP_PORT = int(os.getenv("REPORT_SERVICE_PORT", "8006"))
MONGODB_URL = os.getenv("MONGODB_URL", "mongodb://localhost:27017")

app = FastAPI(title="Report Service", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ==========================================
# KẾT NỐI MONGODB
# ==========================================
client = AsyncIOMotorClient(MONGODB_URL)
test_db = client["TestService"]
student_db = client["StudentService"]

submissions_col = test_db["submissions"]
tests_col = test_db["tests"]
students_col = student_db["students"]


# ==========================================
# HELPER FUNCTIONS
# ==========================================

async def get_all_submissions_raw(test_id: Optional[str] = None, student_code: Optional[str] = None):
    """Lấy toàn bộ submissions, có thể lọc theo test_id hoặc student_code"""
    query = {}
    if test_id:
        query["test_id"] = test_id
    if student_code:
        query["student_code"] = student_code

    docs = await submissions_col.find(query).sort("submitted_at", -1).to_list(5000)
    for d in docs:
        d["_id"] = str(d["_id"])
        if isinstance(d.get("submitted_at"), datetime):
            d["submitted_at"] = d["submitted_at"].isoformat()
    return docs


async def get_test_name_map():
    """Tạo dict: test_id -> {test_code, name}"""
    tests = await tests_col.find().to_list(1000)
    mapping = {}
    for t in tests:
        tid = str(t["_id"])
        mapping[tid] = {
            "test_code": t.get("test_code", ""),
            "name": t.get("name", ""),
            "time": t.get("time", 0),
        }
    return mapping


async def get_student_name_map():
    """Tạo dict: student_code -> student_name"""
    students = await students_col.find().to_list(5000)
    mapping = {}
    for s in students:
        code = s.get("student_code", "")
        if code:
            mapping[code] = s.get("student_name") or s.get("name") or ""
    return mapping


def build_excel(df_main: pd.DataFrame, df_stats: pd.DataFrame, sheet_main: str = "Bảng điểm") -> io.BytesIO:
    """
    Tạo file Excel 2 sheet:
      - Sheet 1: Bảng điểm chi tiết
      - Sheet 2: Thống kê tổng hợp
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # ---- Sheet 1: Chi tiết ----
        df_main.to_excel(writer, sheet_name=sheet_main, index=False)
        ws1 = writer.sheets[sheet_main]

        # Style header
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        header_fill = PatternFill("solid", fgColor="4472C4")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws1[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        # Auto-width columns
        for col in ws1.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws1.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        # Color pass/fail rows
        pass_fill = PatternFill("solid", fgColor="E2EFDA")  # xanh nhạt
        fail_fill = PatternFill("solid", fgColor="FFE0E0")  # đỏ nhạt

        score_col_idx = None
        for idx, col_name in enumerate(df_main.columns, 1):
            if col_name == "Điểm":
                score_col_idx = idx
                break

        if score_col_idx:
            for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
                score_cell = row[score_col_idx - 1]
                try:
                    score = float(score_cell.value or 0)
                    fill = pass_fill if score >= 5 else fail_fill
                    for cell in row:
                        cell.fill = fill
                        cell.border = border
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                except (ValueError, TypeError):
                    pass

        ws1.freeze_panes = "A2"

        # ---- Sheet 2: Thống kê ----
        df_stats.to_excel(writer, sheet_name="Thống kê", index=False)
        ws2 = writer.sheets["Thống kê"]
        stat_fill = PatternFill("solid", fgColor="70AD47")
        stat_font = Font(color="FFFFFF", bold=True)
        for cell in ws2[1]:
            cell.fill = stat_fill
            cell.font = stat_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for col in ws2.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=12)
            ws2.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
        for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

    output.seek(0)
    return output


# ==========================================
# ENDPOINTS
# ==========================================

@app.get("/report_service/health")
async def health():
    return {"status": "ok", "service": "Report Service", "port": FAPP_PORT}


@app.get("/report_service/export/all")
async def export_all_results():
    """
    Xuất toàn bộ bảng điểm tất cả bài thi ra file Excel.
    Dùng bởi Admin → Xuất báo cáo tổng hợp.
    """
    subs = await get_all_submissions_raw()
    test_map = await get_test_name_map()
    student_map = await get_student_name_map()

    if not subs:
        # Trả về file rỗng có header
        df_main = pd.DataFrame(columns=["STT", "Mã SV", "Họ tên", "Mã đề", "Tên bài thi", "Điểm", "Số câu đúng", "Tổng câu", "Kết quả", "Thời gian nộp"])
        df_stats = pd.DataFrame(columns=["Chỉ số", "Giá trị"])
        output = build_excel(df_main, df_stats, "Bảng điểm tổng hợp")
        filename = f"BangDiem_TongHop_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
        )

    rows = []
    for i, s in enumerate(subs, 1):
        tid = s.get("test_id", "")
        test_info = test_map.get(tid, {})
        sc = s.get("student_code", "")
        rows.append({
            "STT": i,
            "Mã SV": sc,
            "Họ tên": student_map.get(sc, ""),
            "Mã đề": test_info.get("test_code", tid),
            "Tên bài thi": test_info.get("name", ""),
            "Điểm": s.get("score", 0),
            "Số câu đúng": s.get("correct_count", 0),
            "Tổng câu": s.get("total_questions", 0),
            "Kết quả": "Đạt ✓" if (s.get("score") or 0) >= 5 else "Chưa đạt ✗",
            "Thời gian nộp": s.get("submitted_at", ""),
        })

    df_main = pd.DataFrame(rows)

    # Thống kê tổng hợp
    scores = [s.get("score", 0) for s in subs]
    pass_count = sum(1 for sc in scores if sc >= 5)
    fail_count = len(scores) - pass_count

    stats_rows = [
        {"Chỉ số": "Tổng số lượt nộp bài", "Giá trị": len(subs)},
        {"Chỉ số": "Số lượt đạt (≥5 điểm)", "Giá trị": pass_count},
        {"Chỉ số": "Số lượt chưa đạt (<5 điểm)", "Giá trị": fail_count},
        {"Chỉ số": "Tỉ lệ đạt (%)", "Giá trị": f"{pass_count / len(scores) * 100:.1f}%" if scores else "0%"},
        {"Chỉ số": "Điểm trung bình", "Giá trị": f"{sum(scores) / len(scores):.2f}" if scores else "0"},
        {"Chỉ số": "Điểm cao nhất", "Giá trị": max(scores) if scores else 0},
        {"Chỉ số": "Điểm thấp nhất", "Giá trị": min(scores) if scores else 0},
        {"Chỉ số": "Số sinh viên thi", "Giá trị": len(set(s.get("student_code") for s in subs))},
        {"Chỉ số": "Số bài thi khác nhau", "Giá trị": len(set(s.get("test_id") for s in subs))},
        {"Chỉ số": "Thời gian xuất báo cáo", "Giá trị": datetime.now().strftime("%d/%m/%Y %H:%M:%S")},
    ]
    df_stats = pd.DataFrame(stats_rows)

    output = build_excel(df_main, df_stats, "Bảng điểm tổng hợp")
    filename = f"BangDiem_TongHop_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )


@app.get("/report_service/export/by_test/{test_id}")
async def export_by_test(test_id: str):
    """
    Xuất bảng điểm của một bài thi cụ thể (theo test_id).
    """
    subs = await get_all_submissions_raw(test_id=test_id)
    test_map = await get_test_name_map()
    student_map = await get_student_name_map()
    test_info = test_map.get(test_id, {})
    test_code = test_info.get("test_code", test_id)
    test_name = test_info.get("name", "")

    rows = []
    for i, s in enumerate(subs, 1):
        sc = s.get("student_code", "")
        rows.append({
            "STT": i,
            "Mã SV": sc,
            "Họ tên": student_map.get(sc, ""),
            "Điểm": s.get("score", 0),
            "Số câu đúng": s.get("correct_count", 0),
            "Tổng câu": s.get("total_questions", 0),
            "Kết quả": "Đạt ✓" if (s.get("score") or 0) >= 5 else "Chưa đạt ✗",
            "Thời gian nộp": s.get("submitted_at", ""),
        })

    df_main = pd.DataFrame(rows) if rows else pd.DataFrame(
        columns=["STT", "Mã SV", "Họ tên", "Điểm", "Số câu đúng", "Tổng câu", "Kết quả", "Thời gian nộp"]
    )

    scores = [s.get("score", 0) for s in subs]
    pass_count = sum(1 for sc in scores if sc >= 5)
    stats_rows = [
        {"Chỉ số": "Tên bài thi", "Giá trị": test_name},
        {"Chỉ số": "Mã đề", "Giá trị": test_code},
        {"Chỉ số": "Tổng số lượt nộp", "Giá trị": len(subs)},
        {"Chỉ số": "Số lượt đạt (≥5)", "Giá trị": pass_count},
        {"Chỉ số": "Số lượt chưa đạt (<5)", "Giá trị": len(subs) - pass_count},
        {"Chỉ số": "Tỉ lệ đạt (%)", "Giá trị": f"{pass_count / len(scores) * 100:.1f}%" if scores else "0%"},
        {"Chỉ số": "Điểm trung bình", "Giá trị": f"{sum(scores) / len(scores):.2f}" if scores else "0"},
        {"Chỉ số": "Điểm cao nhất", "Giá trị": max(scores) if scores else 0},
        {"Chỉ số": "Điểm thấp nhất", "Giá trị": min(scores) if scores else 0},
        {"Chỉ số": "Thời gian xuất", "Giá trị": datetime.now().strftime("%d/%m/%Y %H:%M:%S")},
    ]
    df_stats = pd.DataFrame(stats_rows)

    output = build_excel(df_main, df_stats, f"BĐ_{test_code}")
    safe_code = test_code.replace("/", "_").replace("\\", "_")
    filename = f"BangDiem_{safe_code}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )


@app.get("/report_service/export/by_student/{student_code}")
async def export_by_student(student_code: str):
    """
    Xuất lịch sử thi của một sinh viên ra Excel.
    """
    subs = await get_all_submissions_raw(student_code=student_code)
    test_map = await get_test_name_map()
    student_map = await get_student_name_map()
    student_name = student_map.get(student_code, "")

    rows = []
    for i, s in enumerate(subs, 1):
        tid = s.get("test_id", "")
        test_info = test_map.get(tid, {})
        rows.append({
            "STT": i,
            "Mã đề": test_info.get("test_code", tid),
            "Tên bài thi": test_info.get("name", ""),
            "Điểm": s.get("score", 0),
            "Số câu đúng": s.get("correct_count", 0),
            "Tổng câu": s.get("total_questions", 0),
            "Kết quả": "Đạt ✓" if (s.get("score") or 0) >= 5 else "Chưa đạt ✗",
            "Thời gian nộp": s.get("submitted_at", ""),
        })

    df_main = pd.DataFrame(rows) if rows else pd.DataFrame(
        columns=["STT", "Mã đề", "Tên bài thi", "Điểm", "Số câu đúng", "Tổng câu", "Kết quả", "Thời gian nộp"]
    )

    scores = [s.get("score", 0) for s in subs]
    pass_count = sum(1 for sc in scores if sc >= 5)
    stats_rows = [
        {"Chỉ số": "Mã sinh viên", "Giá trị": student_code},
        {"Chỉ số": "Họ và tên", "Giá trị": student_name},
        {"Chỉ số": "Tổng số lần thi", "Giá trị": len(subs)},
        {"Chỉ số": "Số lần đạt (≥5)", "Giá trị": pass_count},
        {"Chỉ số": "Số lần chưa đạt (<5)", "Giá trị": len(subs) - pass_count},
        {"Chỉ số": "Điểm trung bình", "Giá trị": f"{sum(scores) / len(scores):.2f}" if scores else "0"},
        {"Chỉ số": "Điểm cao nhất", "Giá trị": max(scores) if scores else 0},
        {"Chỉ số": "Điểm thấp nhất", "Giá trị": min(scores) if scores else 0},
        {"Chỉ số": "Thời gian xuất", "Giá trị": datetime.now().strftime("%d/%m/%Y %H:%M:%S")},
    ]
    df_stats = pd.DataFrame(stats_rows)

    output = build_excel(df_main, df_stats, f"LS_{student_code}")
    filename = f"LichSuThi_{student_code}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )


@app.get("/report_service/stats")
async def get_stats():
    """
    Lấy số liệu thống kê tổng hợp (JSON) để hiển thị trên dashboard.
    """
    subs = await get_all_submissions_raw()
    test_map = await get_test_name_map()
    student_map = await get_student_name_map()

    if not subs:
        return {
            "status": "success",
            "data": {
                "total_submissions": 0,
                "pass_count": 0,
                "fail_count": 0,
                "pass_rate": 0,
                "average_score": 0,
                "max_score": 0,
                "min_score": 0,
                "total_students": 0,
                "by_test": []
            }
        }

    scores = [s.get("score", 0) for s in subs]
    pass_count = sum(1 for sc in scores if sc >= 5)

    # Thống kê theo từng bài thi
    by_test = {}
    for s in subs:
        tid = s.get("test_id", "")
        if tid not in by_test:
            ti = test_map.get(tid, {})
            by_test[tid] = {
                "test_id": tid,
                "test_code": ti.get("test_code", tid),
                "test_name": ti.get("name", ""),
                "submissions": 0,
                "pass": 0,
                "fail": 0,
                "scores": []
            }
        by_test[tid]["submissions"] += 1
        sc = s.get("score", 0)
        by_test[tid]["scores"].append(sc)
        if sc >= 5:
            by_test[tid]["pass"] += 1
        else:
            by_test[tid]["fail"] += 1

    by_test_list = []
    for tid, info in by_test.items():
        sc_list = info.pop("scores")
        info["average_score"] = round(sum(sc_list) / len(sc_list), 2) if sc_list else 0
        info["pass_rate"] = round(info["pass"] / info["submissions"] * 100, 1) if info["submissions"] else 0
        by_test_list.append(info)

    return {
        "status": "success",
        "data": {
            "total_submissions": len(subs),
            "pass_count": pass_count,
            "fail_count": len(subs) - pass_count,
            "pass_rate": round(pass_count / len(scores) * 100, 1) if scores else 0,
            "average_score": round(sum(scores) / len(scores), 2) if scores else 0,
            "max_score": max(scores) if scores else 0,
            "min_score": min(scores) if scores else 0,
            "total_students": len(set(s.get("student_code") for s in subs)),
            "by_test": by_test_list
        }
    }


# ==========================================
# MAIN
# ==========================================
if __name__ == "__main__":
    uvicorn.run("main:app", host="127.0.0.1", port=FAPP_PORT, reload=True)
